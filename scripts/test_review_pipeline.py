# review_pipeline.py의 파일명 파싱·그룹핑·리스크 계산·엑셀 생성 로직을 검증하는 테스트
import openpyxl
import tempfile
import unittest
from pathlib import Path

from review_pipeline import compute_risk, group_files_by_product, load_star_dataframe, main, parse_filename, risk_score_label, sample_size_label, write_product_excel


class ParseFilenameTests(unittest.TestCase):
    def test_simple_prefix_and_star(self):
        self.assertEqual(parse_filename("shoe3.csv"), ("shoe", 3))

    def test_ignores_parenthetical_part_number(self):
        self.assertEqual(parse_filename("shoe1 (1).csv"), ("shoe", 1))
        self.assertEqual(parse_filename("shoe1 (2).csv"), ("shoe", 1))

    def test_ignores_dash_part_number(self):
        self.assertEqual(parse_filename("shoe2-1.csv"), ("shoe", 2))

    def test_korean_prefix(self):
        self.assertEqual(parse_filename("손풍기4.csv"), ("손풍기", 4))

    def test_star_out_of_range_returns_none(self):
        self.assertIsNone(parse_filename("shoe6.csv"))
        self.assertIsNone(parse_filename("shoe0.csv"))

    def test_no_digit_returns_none(self):
        self.assertIsNone(parse_filename("shoe.csv"))


class GroupFilesByProductTests(unittest.TestCase):
    def test_groups_by_prefix_and_star(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            (tmp_path / "shoe1 (1).csv").write_text("a,b\n1,2\n", encoding="utf-8")
            (tmp_path / "shoe1 (2).csv").write_text("a,b\n3,4\n", encoding="utf-8")
            (tmp_path / "shoe2.csv").write_text("a,b\n5,6\n", encoding="utf-8")
            (tmp_path / "fan3.csv").write_text("a,b\n7,8\n", encoding="utf-8")
            (tmp_path / "notes.txt").write_text("무시되어야 함", encoding="utf-8")

            groups = group_files_by_product(tmp_path)

            self.assertEqual(set(groups.keys()), {"shoe", "fan"})
            self.assertEqual(len(groups["shoe"][1]), 2)
            self.assertEqual(len(groups["shoe"][2]), 1)
            self.assertEqual(len(groups["fan"][3]), 1)


class RiskScoreLabelTests(unittest.TestCase):
    def test_boundaries(self):
        self.assertEqual(risk_score_label(0.0), "리스크 낮음")
        self.assertEqual(risk_score_label(0.199), "리스크 낮음")
        self.assertEqual(risk_score_label(0.20), "리스크 보통")
        self.assertEqual(risk_score_label(0.399), "리스크 보통")
        self.assertEqual(risk_score_label(0.40), "리스크 높음")
        self.assertEqual(risk_score_label(0.599), "리스크 높음")
        self.assertEqual(risk_score_label(0.60), "리스크 매우 높음")
        self.assertEqual(risk_score_label(1.0), "리스크 매우 높음")


class SampleSizeLabelTests(unittest.TestCase):
    def test_boundaries(self):
        self.assertEqual(sample_size_label(0), "표본 매우 부족")
        self.assertEqual(sample_size_label(9), "표본 매우 부족")
        self.assertEqual(sample_size_label(10), "표본 부족")
        self.assertEqual(sample_size_label(29), "표본 부족")
        self.assertEqual(sample_size_label(30), "표본 충분")


class ComputeRiskTests(unittest.TestCase):
    def test_all_one_star_gives_max_risk(self):
        result = compute_risk(n1=10, n2=0, n3=0, n4=0)
        self.assertEqual(result["n_total"], 10)
        self.assertAlmostEqual(result["risk_score"], 1.0)
        self.assertEqual(result["risk_label"], "리스크 매우 높음")
        self.assertEqual(result["sample_label"], "표본 부족")

    def test_all_four_star_clamps_to_zero(self):
        result = compute_risk(n1=0, n2=0, n3=0, n4=10)
        self.assertAlmostEqual(result["risk_score"], 0.0)
        self.assertEqual(result["risk_label"], "리스크 낮음")

    def test_zero_reviews_returns_not_computable(self):
        result = compute_risk(n1=0, n2=0, n3=0, n4=0)
        self.assertIsNone(result["risk_score"])
        self.assertEqual(result["risk_label"], "계산 불가 (1~4점 리뷰 없음)")
        self.assertEqual(result["sample_label"], "계산 불가")
        self.assertIsNone(result["p1"])
        self.assertIsNone(result["p2"])
        self.assertIsNone(result["p3"])
        self.assertIsNone(result["p4"])

    def test_mixed_counts(self):
        # n_total=20, p1=0.5,p2=0.25,p3=0.15,p4=0.10
        # raw = (4*0.5 + 3*0.25 + 2*0.15 - 1*0.10) / 4 = (2 + 0.75 + 0.3 - 0.1)/4 = 2.95/4 = 0.7375
        result = compute_risk(n1=10, n2=5, n3=3, n4=2)
        self.assertEqual(result["n_total"], 20)
        self.assertAlmostEqual(result["risk_score"], 0.7375)
        self.assertEqual(result["risk_label"], "리스크 매우 높음")
        self.assertEqual(result["sample_label"], "표본 부족")
        self.assertAlmostEqual(result["p1"], 0.5)
        self.assertAlmostEqual(result["p2"], 0.25)
        self.assertAlmostEqual(result["p3"], 0.15)
        self.assertAlmostEqual(result["p4"], 0.10)


class LoadStarDataframeTests(unittest.TestCase):
    def test_adds_qual_columns_when_requested(self):
        with tempfile.TemporaryDirectory() as tmp:
            csv_path = Path(tmp) / "shoe1.csv"
            csv_path.write_text("review_text,date\n소리가 심해요,2025.01.01\n", encoding="utf-8")

            df = load_star_dataframe([csv_path], add_qual_columns=True)

            self.assertEqual(len(df), 1)
            self.assertIn("핵심_보조_태그", df.columns)
            self.assertIn("근거_원문_인용", df.columns)
            self.assertIn("비고", df.columns)

    def test_skips_qual_columns_when_not_requested(self):
        with tempfile.TemporaryDirectory() as tmp:
            csv_path = Path(tmp) / "shoe5.csv"
            csv_path.write_text("review_text,date\n완전 만족해요,2025.01.02\n", encoding="utf-8")

            df = load_star_dataframe([csv_path], add_qual_columns=False)

            self.assertNotIn("핵심_보조_태그", df.columns)


class WriteProductExcelTests(unittest.TestCase):
    def test_creates_workbook_with_expected_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            star1 = tmp_path / "shoe1.csv"
            star1.write_text("review_text,date\n소리 나요,2025.01.01\n실밥 터짐,2025.01.02\n", encoding="utf-8")
            star4 = tmp_path / "shoe4.csv"
            star4.write_text("review_text,date\n대체로 만족,2025.02.01\n", encoding="utf-8")

            star_files = {1: [star1], 4: [star4]}
            out_path = write_product_excel("shoe", star_files, tmp_path)

            self.assertTrue(out_path.exists())
            self.assertEqual(out_path.name, "shoe_review_analysis.xlsx")

            wb = openpyxl.load_workbook(out_path)
            self.assertEqual(
                set(wb.sheetnames), {"요약", "1점", "4점"}
            )

            sheet1 = wb["1점"]
            header = [cell.value for cell in sheet1[1]]
            self.assertIn("핵심_보조_태그", header)
            self.assertEqual(sheet1.max_row, 3)  # 헤더 1행 + 데이터 2행

            summary = wb["요약"]
            summary_values = [cell.value for row in summary.iter_rows() for cell in row]
            self.assertIn("리스크 점수", summary_values)
            self.assertIn("1점 비율", summary_values)
            self.assertIn("66.7%", summary_values)  # star1: 2 of 3 (1~4점 합계) reviews
            self.assertIn("4점 비율", summary_values)
            self.assertIn("33.3%", summary_values)  # star4: 1 of 3 reviews


class MainTests(unittest.TestCase):
    def test_processes_all_products_in_folder(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            (tmp_path / "shoe1.csv").write_text(
                "review_text,date\n별로예요,2025.01.01\n", encoding="utf-8"
            )
            (tmp_path / "fan4.csv").write_text(
                "review_text,date\n괜찮아요,2025.01.02\n", encoding="utf-8"
            )

            outputs = main(str(tmp_path))

            names = {p.name for p in outputs}
            self.assertEqual(names, {"shoe_review_analysis.xlsx", "fan_review_analysis.xlsx"})
            for p in outputs:
                self.assertTrue(p.exists())


if __name__ == "__main__":
    unittest.main()
